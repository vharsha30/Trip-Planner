from docx import Document


def export_itinerary(itinerary, filename):

    doc = Document()

    doc.add_heading("Trip Itinerary", level=1)

    for day, places in itinerary.items():

        doc.add_heading(day, level=2)

        for place in places:

            doc.add_paragraph(place)

    doc.save(filename)